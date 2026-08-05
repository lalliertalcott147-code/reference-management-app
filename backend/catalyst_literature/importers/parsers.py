from __future__ import annotations

import io
import re
from collections.abc import Iterable
from typing import Any

from openpyxl import load_workbook

from ..search.models import PaperRecord
from .models import ImportIssue, ImportValidationError, ParsedImport
from .validation import decode_text, detect_format

TAGGED_LINE = re.compile(r"^([A-Z0-9]{2}) (.*)$")
RIS_LINE = re.compile(r"^([A-Z0-9]{2})  - ?(.*)$")


def _year(value: str | None) -> int | None:
    if not value:
        return None
    match = re.search(r"\b(1[6-9]\d{2}|20\d{2}|21\d{2})\b", value)
    return int(match.group(1)) if match else None


def _first(fields: dict[str, list[str]], *names: str) -> str | None:
    for name in names:
        values = fields.get(name, [])
        if values and values[0].strip():
            return values[0].strip()
    return None


def _record_from_fields(
    fields: dict[str, list[str]],
    *,
    mapping: dict[str, tuple[str, ...]],
    record_number: int,
) -> tuple[PaperRecord | None, ImportIssue | None]:
    title = _first(fields, *mapping["title"])
    if not title:
        return None, ImportIssue(record_number, "title", "Record has no title and was skipped")
    authors: list[str] = []
    for key in mapping["authors"]:
        for value in fields.get(key, []):
            authors.extend(part.strip() for part in value.split(";") if part.strip())
    uid = _first(fields, *mapping["uid"])
    doi = _first(fields, *mapping["doi"])
    source_id = uid or doi or f"row-{record_number}"
    return (
        PaperRecord(
            source="wos_export",
            source_id=source_id,
            wos_uid=uid,
            title=title,
            doi=doi,
            authors=tuple(authors),
            journal=_first(fields, *mapping["journal"]),
            year=_year(_first(fields, *mapping["year"])),
            abstract=_first(fields, *mapping["abstract"]),
            url=_first(fields, *mapping["url"]),
            raw={key: values for key, values in fields.items()},
        ),
        None,
    )


def parse_plain_text(data: bytes) -> ParsedImport:
    text = decode_text(data)
    records: list[PaperRecord] = []
    issues: list[ImportIssue] = []
    current: dict[str, list[str]] = {}
    current_tag: str | None = None

    def finish() -> None:
        nonlocal current, current_tag
        if not current:
            return
        record, issue = _record_from_fields(
            current,
            mapping={
                "title": ("TI",),
                "authors": ("AU", "AF"),
                "journal": ("SO",),
                "year": ("PY",),
                "doi": ("DI",),
                "uid": ("UT",),
                "abstract": ("AB",),
                "url": ("UR",),
            },
            record_number=len(records) + len(issues) + 1,
        )
        if record:
            records.append(record)
        if issue:
            issues.append(issue)
        current = {}
        current_tag = None

    for raw_line in text.splitlines():
        if raw_line == "ER":
            finish()
            continue
        match = TAGGED_LINE.match(raw_line)
        if match:
            tag, value = match.groups()
            if tag in {"FN", "VR", "EF"} and not current:
                continue
            current.setdefault(tag, []).append(value.strip())
            current_tag = tag
        elif raw_line.startswith("   ") and current_tag and current.get(current_tag):
            current[current_tag][-1] += " " + raw_line.strip()
    finish()
    return ParsedImport("wos_plain_text", tuple(records), tuple(issues))


def parse_ris(data: bytes) -> ParsedImport:
    text = decode_text(data)
    records: list[PaperRecord] = []
    issues: list[ImportIssue] = []
    current: dict[str, list[str]] = {}
    current_tag: str | None = None
    for raw_line in text.splitlines():
        match = RIS_LINE.match(raw_line)
        if match:
            tag, value = match.groups()
            if tag == "ER":
                record, issue = _record_from_fields(
                    current,
                    mapping={
                        "title": ("TI", "T1"),
                        "authors": ("AU", "A1"),
                        "journal": ("JO", "JF", "T2"),
                        "year": ("PY", "Y1"),
                        "doi": ("DO",),
                        "uid": ("AN", "UT"),
                        "abstract": ("AB", "N2"),
                        "url": ("UR",),
                    },
                    record_number=len(records) + len(issues) + 1,
                )
                if record:
                    records.append(record)
                if issue:
                    issues.append(issue)
                current = {}
                current_tag = None
            else:
                current.setdefault(tag, []).append(value.strip())
                current_tag = tag
        elif raw_line.startswith("  ") and current_tag and current.get(current_tag):
            current[current_tag][-1] += " " + raw_line.strip()
    if current:
        issues.append(ImportIssue(None, None, "The final RIS record has no ER terminator"))
    return ParsedImport("ris", tuple(records), tuple(issues))


EXCEL_HEADERS: dict[str, tuple[str, ...]] = {
    "title": ("article title", "title", "document title"),
    "authors": ("authors", "author full names", "author"),
    "journal": ("source title", "journal", "publication name"),
    "year": ("publication year", "year published", "year"),
    "doi": ("doi", "digital object identifier"),
    "uid": ("accession number", "ut (unique wos id)", "wos uid", "ut"),
    "abstract": ("abstract", "ab"),
    "url": ("wos url", "url"),
}


def _string(value: Any) -> str:
    return "" if value is None else str(value).strip()


def parse_xlsx(data: bytes) -> ParsedImport:
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    try:
        if not workbook.worksheets:
            raise ImportValidationError("The workbook contains no worksheet")
        sheet = workbook.worksheets[0]
        rows: Iterable[tuple[Any, ...]] = sheet.iter_rows(values_only=False)
        try:
            header_cells = next(iter(rows))
        except StopIteration as error:
            raise ImportValidationError("The workbook contains no header row") from error
        if len(header_cells) > 200:
            raise ImportValidationError("The workbook has too many columns")
        header_lookup = {
            _string(cell.value).casefold(): index for index, cell in enumerate(header_cells)
        }
        column_for: dict[str, int] = {}
        for field, aliases in EXCEL_HEADERS.items():
            for alias in aliases:
                if alias in header_lookup:
                    column_for[field] = header_lookup[alias]
                    break
        if "title" not in column_for:
            raise ImportValidationError("The workbook has no recognized title column")
        records: list[PaperRecord] = []
        issues: list[ImportIssue] = []
        for row_number, cells in enumerate(rows, start=2):
            if row_number > 10_001:
                raise ImportValidationError("The workbook exceeds the 10,000 record limit")
            if any(cell.data_type == "f" for cell in cells):
                raise ImportValidationError(
                    f"Formula cells are not accepted in import files (row {row_number})"
                )
            values = tuple(_string(cell.value) for cell in cells)
            if not any(values):
                continue
            fields = {
                field: [values[index]]
                for field, index in column_for.items()
                if index < len(values) and values[index]
            }
            record, issue = _record_from_fields(
                fields,
                mapping={key: (key,) for key in EXCEL_HEADERS},
                record_number=row_number - 1,
            )
            if record:
                records.append(record)
            if issue:
                issues.append(issue)
        return ParsedImport("xlsx", tuple(records), tuple(issues))
    finally:
        workbook.close()


def parse_wos_export(filename: str, data: bytes) -> ParsedImport:
    detected = detect_format(filename, data)
    if detected == "wos_plain_text":
        return parse_plain_text(data)
    if detected == "ris":
        return parse_ris(data)
    return parse_xlsx(data)
